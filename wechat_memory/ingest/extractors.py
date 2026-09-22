"""Per-type text extraction (DESIGN.md §6.4).

Each function returns extracted text for one attachment file; the pipeline
decides which to call based on the detected subtype.  All are pure functions
over file paths so they are easy to test.
"""

from __future__ import annotations

import logging
import zipfile
from pathlib import Path

logger = logging.getLogger(__name__)

# Text-ish files worth reading inside a source zip (small, capped).
_ZIP_READABLE = {".md", ".txt", ".rst", ".adoc", ".markdown"}
_ZIP_READ_LIMIT = 4  # max files to read per zip
_ZIP_FILE_CAP = 512 * 1024  # 512KB per file


def extract_pdf(path: Path) -> str:
    """Extract text from a PDF via PyMuPDF."""
    try:
        import pymupdf  # prefer new name; fallback to fitz
    except ImportError:  # pragma: no cover
        import fitz as pymupdf
    pages = []
    with pymupdf.open(str(path)) as doc:
        for page in doc:
            pages.append(page.get_text())
    return "\n".join(pages).strip()


def extract_docx(path: Path) -> str:
    """Extract paragraphs from a .docx via python-docx."""
    import docx

    d = docx.Document(str(path))
    parts = [p.text for p in d.paragraphs if p.text.strip()]
    # Tables carry meaningful content too.
    for table in d.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells]
            parts.append(" | ".join(c for c in cells if c))
    return "\n".join(parts).strip()


def extract_xlsx(path: Path) -> str:
    """Extract each sheet's cells via openpyxl, preserving row structure."""
    import openpyxl

    wb = openpyxl.load_workbook(str(path), data_only=True, read_only=True)
    out = []
    for ws in wb.worksheets:
        out.append(f"# Sheet: {ws.title}")
        for row in ws.iter_rows(values_only=True):
            cells = [str(c).strip() for c in row if c is not None and str(c).strip()]
            if cells:
                out.append(" | ".join(cells))
    wb.close()
    return "\n".join(out).strip()


def extract_zip_text(path: Path) -> str:
    """Shallow-check a zip for readable docs, return their text (capped)."""
    try:
        zf = zipfile.ZipFile(str(path))
    except zipfile.BadZipFile:
        return ""
    out = []
    count = 0
    with zf:
        for info in zf.infolist():
            if info.is_dir() or count >= _ZIP_READ_LIMIT:
                continue
            suffix = Path(info.filename).suffix.lower()
            if suffix not in _ZIP_READABLE or info.file_size > _ZIP_FILE_CAP:
                continue
            try:
                text = zf.read(info).decode("utf-8", errors="replace")
            except Exception:  # noqa: BLE001 - a bad entry shouldn't kill the zip
                continue
            out.append(f"## {info.filename}\n{text}")
            count += 1
    return "\n\n".join(out).strip()


def ocr_image(path: Path) -> str:
    """OCR an image via RapidOCR. Returns '' when no text found."""
    from rapidocr_onnxruntime import RapidOCR

    ocr = RapidOCR()
    result, _ = ocr(str(path))
    if not result:
        return ""
    lines = [r[1] for r in result]
    return "\n".join(lines).strip()
